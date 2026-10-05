import datetime
from openpyxl import Workbook, load_workbook
import os
import sys
from pathlib import Path

def get_save_path(location, date_str):
    filename = f"{location}_{date_str}.xlsx"

    # On iOS (Pyto), use sandboxed Documents folder
    if sys.platform == "ios":
        documents_dir = Path.home() / "Documents"
        return documents_dir / filename

    # On macOS or other platforms, save in current directory
    return Path(filename)

def main():
    # Ask for location
    location = input("Enter Location: ").strip()

    # Build filename with date
    date_str = datetime.datetime.now().strftime("%m%d%y")
    filepath = get_save_path(location, date_str)

    # Load or create workbook
    if filepath.exists():
        wb = load_workbook(filepath)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Timestamp", "Barcode"])  # header row

    print(f"\n✅ Scanning barcodes for location '{location}'")
    print(f"📁 Saving to: {filepath}")
    print("📦 Scan barcodes (Ctrl+C to stop)\n")

    try:
        while True:
            barcode = input().strip()
            if barcode:
                timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                ws.append([timestamp, barcode])
                wb.save(filepath)
                print(f"✔ Saved: {barcode}")
    except KeyboardInterrupt:
        print("\n🛑 Stopped. All scans saved.")

if __name__ == "__main__":
    main()